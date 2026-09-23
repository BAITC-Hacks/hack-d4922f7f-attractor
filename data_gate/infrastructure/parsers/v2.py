"""Explicit V2 source adapters; never evaluate formulas or fetch external links."""

from __future__ import annotations

import csv
import io
import zipfile
from datetime import date, datetime
from typing import Any

from data_gate.application.ports import ParsedSource
from data_gate.domain.errors import SourceParseError
from data_gate.domain.v2_quality import OBSERVATIONS, V2_CITY
from data_gate.infrastructure.parsers.city_json import load_json

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 100_000
MAX_COLUMNS = 100
CITY_FIELDS = (
    "schemaVersion", "mode", "status", "currency", "basePeriod", "startSimTime", "districts",
    "weather", "finance", "parameters", "provenance", "units", "assumptions",
)


def _bounded(content: bytes) -> None:
    if not content or len(content) > MAX_BYTES:
        raise SourceParseError("Source must contain 1 byte to 10 MiB", field="content")


class V2CityJsonParser:
    format_id = "v2-city-json"

    def parse(self, content: bytes) -> ParsedSource:
        _bounded(content)
        raw = load_json(content)
        if not isinstance(raw, dict):
            raise SourceParseError("Expected a V2 city JSON object")
        payload = {key: raw[key] for key in CITY_FIELDS if key in raw}
        return ParsedSource(V2_CITY, payload, tuple(payload), tuple(sorted(set(raw) - set(payload))))


class ObservationsJsonParser:
    format_id = "observations-json-v1"

    def parse(self, content: bytes) -> ParsedSource:
        _bounded(content)
        raw = load_json(content)
        if not isinstance(raw, dict):
            raise SourceParseError("Expected an observations JSON object")
        payload = {"records": raw.get("records")}
        return ParsedSource(OBSERVATIONS, payload, ("records",), tuple(sorted(set(raw) - {"records"})))


def _records(rows: Any, header_row: int) -> ParsedSource:
    headers: list[str] | None = None
    records: list[dict[str, Any]] = []
    for number, row in enumerate(rows, start=1):
        if number > MAX_ROWS + header_row:
            raise SourceParseError("Table exceeds 100000 data rows")
        if number < header_row:
            continue
        if len(row) > MAX_COLUMNS:
            raise SourceParseError("Table exceeds 100 columns")
        if headers is None:
            headers = [str(value).strip() if value is not None else "" for value in row]
            if not headers or not all(headers) or len(set(headers)) != len(headers):
                raise SourceParseError("Headers must be nonempty and unique", field="headers")
            continue
        if all(value is None or value == "" for value in row):
            continue
        if len(row) != len(headers):
            raise SourceParseError(f"Row {number} has a different number of fields", field=f"rows[{number}]")
        record: dict[str, Any] = {}
        for key, cell in zip(headers, row):
            value = cell.isoformat() if isinstance(cell, (date, datetime)) else cell
            if isinstance(value, str):
                value = value.strip()
            if value == "" or value is None:
                value = None
            elif key == "value" and isinstance(value, str):
                try:
                    value = float(value)
                except ValueError:
                    pass  # Stable quality report, not silent coercion to zero.
            elif key == "imputation" and isinstance(value, str):
                value = load_json(value.encode("utf-8"))
            record[key] = value
        records.append(record)
    if headers is None:
        raise SourceParseError("Selected header row does not exist", field="headerRow")
    return ParsedSource(OBSERVATIONS, {"records": records}, tuple(headers))


class CsvParser:
    format_id = "csv-v1"

    def __init__(self, *, header_row: int = 1, delimiter: str = ",") -> None:
        if not 1 <= header_row <= 100 or len(delimiter) != 1:
            raise ValueError("header_row must be 1..100; delimiter must be one character")
        self.header_row = header_row
        self.delimiter = delimiter

    def parse(self, content: bytes) -> ParsedSource:
        _bounded(content)
        try:
            reader = csv.reader(io.StringIO(content.decode("utf-8-sig")), delimiter=self.delimiter, strict=True)
            return _records(reader, self.header_row)
        except (UnicodeDecodeError, csv.Error) as error:
            raise SourceParseError(f"Invalid UTF-8 CSV: {error}") from error


class XlsxParser:
    format_id = "xlsx-v1"

    def __init__(self, *, sheet: str | None = None, header_row: int = 1) -> None:
        if not 1 <= header_row <= 100:
            raise ValueError("header_row must be 1..100")
        self.sheet = sheet
        self.header_row = header_row

    def parse(self, content: bytes) -> ParsedSource:
        _bounded(content)
        try:
            from openpyxl import load_workbook
        except ImportError as error:
            raise SourceParseError("XLSX requires the openpyxl dependency") from error
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if len(archive.infolist()) > 1000 or sum(item.file_size for item in archive.infolist()) > 5 * MAX_BYTES:
                    raise SourceParseError("XLSX expanded size exceeds 50 MiB or 1000 ZIP entries")
                if any("vbaProject" in item.filename for item in archive.infolist()):
                    raise SourceParseError("Macro-enabled workbooks are not accepted")
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                if self.sheet and self.sheet not in workbook.sheetnames:
                    raise SourceParseError("Selected worksheet does not exist", field="sheet")
                worksheet = workbook[self.sheet] if self.sheet else workbook.worksheets[0]
                if (worksheet.max_row or 0) > MAX_ROWS + self.header_row or (worksheet.max_column or 0) > MAX_COLUMNS:
                    raise SourceParseError("Worksheet dimensions exceed 100000 rows or 100 columns")

                def rows():
                    for row in worksheet.iter_rows():
                        if any(cell.data_type == "f" for cell in row):
                            raise SourceParseError("Formula cells must be replaced with explicit values", field="formula")
                        yield [cell.value for cell in row]

                return _records(rows(), self.header_row)
            finally:
                workbook.close()
        except SourceParseError:
            raise
        except (ValueError, KeyError, OSError, IndexError, SyntaxError, zipfile.BadZipFile) as error:
            raise SourceParseError(f"Invalid XLSX workbook: {error}") from error
